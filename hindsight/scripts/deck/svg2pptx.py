#!/usr/bin/env python3
"""Convert matplotlib SVG figures into native, editable PowerPoint shapes.

Every SVG path becomes a DrawingML freeform (custGeom) with its fill, stroke, dash, cap, join and
opacity; every text element becomes a real text box (mathtext tspans become runs, with sub/superscript
as baseline offsets). Each figure is one grouped object on its own slide, scaled to fit, so it can be
moved as a unit or ungrouped and edited piece by piece. Clip paths are not carried over (the figures
keep their content inside the axes).

usage: svg2pptx.py OUT.pptx FIG1.svg [FIG2.svg ...]
"""
from __future__ import annotations

import math
import re
import sys
from pathlib import Path
from xml.sax.saxutils import escape

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

SVG = "{http://www.w3.org/2000/svg}"
XLINK = "{http://www.w3.org/1999/xlink}"
PT = 12700  # EMU per point
SLIDE_W, SLIDE_H, MARGIN = 12192000, 6858000, int(0.3 * 914400)
FONT = "Arial"
A = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
P = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'


# ---------------------------------------------------------------- geometry
def mul(m, n):
    a, b, c, d, e, f = m
    a2, b2, c2, d2, e2, f2 = n
    return (a * a2 + c * b2, b * a2 + d * b2, a * c2 + c * d2, b * c2 + d * d2, a * e2 + c * f2 + e, b * e2 + d * f2 + f)


def apply(m, x, y):
    a, b, c, d, e, f = m
    return a * x + c * y + e, b * x + d * y + f


I = (1, 0, 0, 1, 0, 0)


def parse_transform(t):
    m = I
    for name, args in re.findall(r"(\w+)\s*\(([^)]*)\)", t or ""):
        v = [float(x) for x in re.split(r"[\s,]+", args.strip()) if x]
        if name == "translate":
            n = (1, 0, 0, 1, v[0], v[1] if len(v) > 1 else 0)
        elif name == "scale":
            n = (v[0], 0, 0, v[1] if len(v) > 1 else v[0], 0, 0)
        elif name == "rotate":
            r = math.radians(v[0]); cs, sn = math.cos(r), math.sin(r)
            n = (cs, sn, -sn, cs, 0, 0)
            if len(v) == 3:
                n = mul(mul((1, 0, 0, 1, v[1], v[2]), n), (1, 0, 0, 1, -v[1], -v[2]))
        elif name == "matrix":
            n = tuple(v)
        else:
            continue
        m = mul(m, n)
    return m


def parse_style(s):
    out = {}
    for part in (s or "").split(";"):
        if ":" in part:
            k, v = part.split(":", 1)
            out[k.strip()] = v.strip()
    return out


def parse_path(d):
    toks = re.findall(r"[MLCQZHVmlcqzhv]|-?(?:\d+\.?\d*|\.\d+)(?:e[-+]?\d+)?", d)
    segs, i, cmd, cur, start = [], 0, None, (0.0, 0.0), (0.0, 0.0)
    num = lambda: float(toks[i])
    while i < len(toks):
        t = toks[i]
        if re.match(r"[A-Za-z]", t):
            cmd = t; i += 1
            if cmd in "Zz":
                segs.append(("Z",)); cur = start
            continue
        if cmd in "Mm":
            x, y = float(toks[i]), float(toks[i + 1]); i += 2
            if cmd == "m": x, y = x + cur[0], y + cur[1]
            segs.append(("M", (x, y))); cur = start = (x, y)
            cmd = "L" if cmd == "M" else "l"
        elif cmd in "Ll":
            x, y = float(toks[i]), float(toks[i + 1]); i += 2
            if cmd == "l": x, y = x + cur[0], y + cur[1]
            segs.append(("L", (x, y))); cur = (x, y)
        elif cmd in "Hh":
            x = float(toks[i]); i += 1
            x = x + cur[0] if cmd == "h" else x
            segs.append(("L", (x, cur[1]))); cur = (x, cur[1])
        elif cmd in "Vv":
            y = float(toks[i]); i += 1
            y = y + cur[1] if cmd == "v" else y
            segs.append(("L", (cur[0], y))); cur = (cur[0], y)
        elif cmd in "Cc":
            v = [float(toks[i + k]) for k in range(6)]; i += 6
            if cmd == "c": v = [v[k] + cur[k % 2] for k in range(6)]
            segs.append(("C", (v[0], v[1]), (v[2], v[3]), (v[4], v[5]))); cur = (v[4], v[5])
        elif cmd in "Qq":
            v = [float(toks[i + k]) for k in range(4)]; i += 4
            if cmd == "q": v = [v[k] + cur[k % 2] for k in range(4)]
            q, e = (v[0], v[1]), (v[2], v[3])
            c1 = (cur[0] + 2 / 3 * (q[0] - cur[0]), cur[1] + 2 / 3 * (q[1] - cur[1]))
            c2 = (e[0] + 2 / 3 * (q[0] - e[0]), e[1] + 2 / 3 * (q[1] - e[1]))
            segs.append(("C", c1, c2, e)); cur = e
        else:
            i += 1
    return segs


# ---------------------------------------------------------------- DrawingML
def color(v):
    if not v or v == "none":
        return None
    v = v.strip()
    if v.startswith("#"):
        v = v[1:]
        if len(v) == 3:
            v = "".join(ch * 2 for ch in v)
        return v.upper()
    named = {"black": "000000", "white": "FFFFFF", "red": "FF0000", "blue": "0000FF", "gray": "808080", "grey": "808080"}
    return named.get(v.lower(), "000000")


def solid(hexv, alpha):
    a = "" if alpha >= 0.999 else f'<a:alpha val="{int(round(alpha * 100000))}"/>'
    return f'<a:solidFill><a:srgbClr val="{hexv}">{a}</a:srgbClr></a:solidFill>'


class Builder:
    def __init__(self, scale, ox, oy):
        self.k, self.ox, self.oy, self.nid, self.parts = scale, ox, oy, 10, []
        self.bbox = [math.inf, math.inf, -math.inf, -math.inf]

    def emu(self, x, y):
        return self.ox + x * PT * self.k, self.oy + y * PT * self.k

    def grow(self, x0, y0, x1, y1):
        b = self.bbox
        b[0], b[1], b[2], b[3] = min(b[0], x0), min(b[1], y0), max(b[2], x1), max(b[3], y1)

    def next_id(self):
        self.nid += 1
        return self.nid

    def path(self, segs, ctm, st, name):
        fill = color(st.get("fill", "#000000"))
        stroke = color(st.get("stroke"))
        if fill is None and stroke is None:
            return
        op = float(st.get("opacity", 1))
        fa, sa = op * float(st.get("fill-opacity", 1)), op * float(st.get("stroke-opacity", 1))
        pts = []
        tseg = []
        for s in segs:
            if s[0] == "Z":
                tseg.append(("Z",)); continue
            q = [self.emu(*apply(ctm, *p)) for p in s[1:]]
            pts += q; tseg.append((s[0], *q))
        if not pts:
            return
        x0, y0 = min(p[0] for p in pts), min(p[1] for p in pts)
        x1, y1 = max(p[0] for p in pts), max(p[1] for p in pts)
        w, h = max(int(round(x1 - x0)), 1), max(int(round(y1 - y0)), 1)
        rel = lambda p: f'<a:pt x="{int(round(p[0] - x0))}" y="{int(round(p[1] - y0))}"/>'
        body = []
        for s in tseg:
            if s[0] == "M": body.append(f"<a:moveTo>{rel(s[1])}</a:moveTo>")
            elif s[0] == "L": body.append(f"<a:lnTo>{rel(s[1])}</a:lnTo>")
            elif s[0] == "C": body.append(f"<a:cubicBezTo>{rel(s[1])}{rel(s[2])}{rel(s[3])}</a:cubicBezTo>")
            elif s[0] == "Z": body.append("<a:close/>")
        sw = float(st.get("stroke-width", 1)) * PT * self.k * math.sqrt(abs(ctm[0] * ctm[3] - ctm[1] * ctm[2]))
        if stroke:
            cap = {"butt": "flat", "round": "rnd", "square": "sq"}.get(st.get("stroke-linecap", "butt"), "flat")
            join = {"round": "<a:round/>", "miter": "<a:miter lim=\"800000\"/>", "bevel": "<a:bevel/>"}.get(st.get("stroke-linejoin", "round"), "<a:round/>")
            dash = ""
            da = st.get("stroke-dasharray")
            if da and da != "none":
                vals = [float(v) for v in re.split(r"[\s,]+", da) if v]
                width_pt = max(float(st.get("stroke-width", 1)), 1e-3)
                if len(vals) % 2: vals = vals * 2
                ds = "".join(f'<a:ds d="{int(vals[j] / width_pt * 100000)}" sp="{int(vals[j + 1] / width_pt * 100000)}"/>'
                             for j in range(0, len(vals), 2))
                dash = f"<a:custDash>{ds}</a:custDash>"
            ln = f'<a:ln w="{max(int(sw), 1)}" cap="{cap}">{solid(stroke, sa)}{dash}{join}</a:ln>'
        else:
            ln = "<a:ln><a:noFill/></a:ln>"
        fl = solid(fill, fa) if fill else "<a:noFill/>"
        sid = self.next_id()
        self.parts.append(
            f'<p:sp><p:nvSpPr><p:cNvPr id="{sid}" name="{escape(name)} {sid}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
            f'<p:spPr><a:xfrm><a:off x="{int(round(x0))}" y="{int(round(y0))}"/><a:ext cx="{w}" cy="{h}"/></a:xfrm>'
            f'<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/><a:rect l="0" t="0" r="r" b="b"/>'
            f'<a:pathLst><a:path w="{w}" h="{h}">{"".join(body)}</a:path></a:pathLst></a:custGeom>{fl}{ln}</p:spPr>'
            f'<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:endParaRPr lang="en-US"/></a:p></p:txBody></p:sp>')
        self.grow(x0, y0, x0 + w, y0 + h)

    def text(self, runs, ctm, ax, ay, anchor, name):
        """runs: [(text, size_pt, bold, italic, color, baseline)]; (ax, ay) = anchor point on the baseline (svg units)."""
        txt = "".join(r[0] for r in runs)
        if not txt.strip():
            return
        size = max(r[1] for r in runs)
        k = self.k * math.sqrt(abs(ctm[0] * ctm[3] - ctm[1] * ctm[2]))
        s_emu = size * PT * k
        w = sum(len(r[0]) * r[1] * (0.6 if r[2] else 0.56) for r in runs) * PT * k + s_emu * 0.3
        h = s_emu * 1.25
        x0 = {"start": 0, "middle": -w / 2, "end": -w}.get(anchor, 0)
        cx_l, cy_l = x0 + w / 2, -0.93 * s_emu + h / 2  # box centre relative to the anchor, in rotated text frame
        theta = math.atan2(ctm[1], ctm[0])
        gx, gy = self.emu(*apply(ctm, ax, ay))
        cx = gx + cx_l * math.cos(theta) - cy_l * math.sin(theta)
        cy = gy + cx_l * math.sin(theta) + cy_l * math.cos(theta)
        offx, offy = cx - w / 2, cy - h / 2
        rot = int(round(math.degrees(theta) * 60000)) % 21600000
        algn = {"start": "l", "middle": "ctr", "end": "r"}.get(anchor, "l")
        rr = []
        for t, sz, b, it, col, base in runs:
            attrs = f'lang="en-US" sz="{max(int(round(size * k * 100)), 100)}"'
            if b: attrs += ' b="1"'
            if it: attrs += ' i="1"'
            if base: attrs += f' baseline="{base}"'
            rr.append(f'<a:r><a:rPr {attrs} dirty="0">{solid(col, 1)}<a:latin typeface="{FONT}"/><a:cs typeface="{FONT}"/></a:rPr>'
                      f'<a:t>{escape(t)}</a:t></a:r>')
        sid = self.next_id()
        rot_attr = f' rot="{rot}"' if rot else ""
        self.parts.append(
            f'<p:sp><p:nvSpPr><p:cNvPr id="{sid}" name="{escape(name)} {sid}"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
            f'<p:spPr><a:xfrm{rot_attr}><a:off x="{int(round(offx))}" y="{int(round(offy))}"/><a:ext cx="{int(w)}" cy="{int(h)}"/></a:xfrm>'
            f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
            f'<p:txBody><a:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" anchor="t"><a:noAutofit/></a:bodyPr><a:lstStyle/>'
            f'<a:p><a:pPr algn="{algn}"/>{"".join(rr)}</a:p></p:txBody></p:sp>')
        self.grow(offx, offy, offx + w, offy + h)

    def group(self, name):
        b = [int(v) for v in self.bbox]
        return (f'<p:grpSp {P} {A}><p:nvGrpSpPr><p:cNvPr id="2" name="{escape(name)}"/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
                f'<p:grpSpPr><a:xfrm><a:off x="{b[0]}" y="{b[1]}"/><a:ext cx="{b[2] - b[0]}" cy="{b[3] - b[1]}"/>'
                f'<a:chOff x="{b[0]}" y="{b[1]}"/><a:chExt cx="{b[2] - b[0]}" cy="{b[3] - b[1]}"/></a:xfrm></p:grpSpPr>'
                + "".join(self.parts) + "</p:grpSp>")


# ---------------------------------------------------------------- traversal
def text_runs(el, inherited):
    st = {**inherited, **parse_style(el.get("style"))}
    def fmt(s):
        size = float(re.match(r"[\d.]+", s.get("font-size", "10")).group(0))
        bold = s.get("font-weight", "normal") in ("bold", "700", "800", "900", "600")
        ital = s.get("font-style", "normal") in ("italic", "oblique")
        return size, bold, ital, color(s.get("fill", "#000000")) or "000000"
    tsp = el.findall(f"{SVG}tspan")
    if not tsp:
        size, b, it, col = fmt(st)
        return [((el.text or ""), size, b, it, col, 0)], float(el.get("x", 0)), float(el.get("y", 0)), st.get("text-anchor", "start")
    items = []
    for t in tsp:
        s = {**st, **parse_style(t.get("style"))}
        size, b, it, col = fmt(s)
        items.append((t.text or "", size, b, it, col, float(t.get("x", 0)), float(t.get("y", 0))))
    items.sort(key=lambda i: i[5])  # mathtext emits glyphs grouped by font, not in reading order
    ys = [i[6] for i in items]
    main_y = max(set(ys), key=lambda y: sum(1 for i in items if i[6] == y and i[1] == max(x[1] for x in items)))
    main_size = max(i[1] for i in items if i[6] == main_y)
    runs = []
    for txt, size, b, it, col, x, y in items:
        base = 0
        if size < main_size - 0.01 or abs(y - main_y) > 0.01:
            base = -25000 if y > main_y + 0.01 else (30000 if y < main_y - 0.01 else 0)
        if runs and runs[-1][1:] == (main_size, b, it, col, base):
            runs[-1] = (runs[-1][0] + txt, main_size, b, it, col, base)
        else:
            runs.append((txt, main_size, b, it, col, base))
    return runs, items[0][5], main_y, "start"


def walk(el, ctm, bld, defs, name, inherited):
    tag = el.tag.replace(SVG, "") if isinstance(el.tag, str) else ""
    if tag in ("defs", "clipPath", "metadata", "style", "title", "desc") or not tag:
        return
    m = mul(ctm, parse_transform(el.get("transform")))
    nm = el.get("id") or name
    if tag in ("svg", "g", "a"):
        for ch in el:
            walk(ch, m, bld, defs, nm, inherited)
    elif tag == "path":
        bld.path(parse_path(el.get("d", "")), m, parse_style(el.get("style")), nm)
    elif tag == "rect":
        x, y, w, h = (float(el.get(a, 0)) for a in ("x", "y", "width", "height"))
        bld.path([("M", (x, y)), ("L", (x + w, y)), ("L", (x + w, y + h)), ("L", (x, y + h)), ("Z",)], m, parse_style(el.get("style")), nm)
    elif tag == "use":
        ref = defs.get((el.get(f"{XLINK}href") or el.get("href") or "").lstrip("#"))
        if ref is None:
            return
        mu = mul(m, (1, 0, 0, 1, float(el.get("x", 0)), float(el.get("y", 0))))
        st = {**parse_style(ref.get("style")), **parse_style(el.get("style"))}
        rtag = ref.tag.replace(SVG, "")
        if rtag == "path":
            bld.path(parse_path(ref.get("d", "")), mul(mu, parse_transform(ref.get("transform"))), st, nm)
    elif tag == "text":
        runs, x, y, anchor = text_runs(el, inherited)
        bld.text(runs, m, x, y, anchor, nm)


def convert(svg_path, prs, title):
    root = etree.parse(str(svg_path)).getroot()
    vb = [float(v) for v in root.get("viewBox").split()]
    fw, fh = vb[2], vb[3]
    k = min((SLIDE_W - 2 * MARGIN) / (fw * PT), (SLIDE_H - 2 * MARGIN) / (fh * PT))
    ox = (SLIDE_W - fw * PT * k) / 2 - vb[0] * PT * k
    oy = (SLIDE_H - fh * PT * k) / 2 - vb[1] * PT * k
    defs = {e.get("id"): e for e in root.iter() if e.get("id") and isinstance(e.tag, str) and e.tag.endswith("path")}
    bld = Builder(k, ox, oy)
    fig = root.find(f"{SVG}g")
    first = True
    for ch in fig:
        if first and ch.get("id") == "patch_1":  # figure background: dropped so the figure sits on any slide colour
            first = False
            continue
        first = False
        walk(ch, I, bld, defs, ch.get("id") or "shape", {})
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    grp = etree.fromstring(bld.group(title))
    slide.shapes._spTree.append(grp)
    # renumber ids uniquely within the slide
    nid = 2
    for c in slide.shapes._spTree.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr"):
        c.set("id", str(nid)); nid += 1
    slide.notes_slide.notes_text_frame.text = title
    return len(bld.parts)


def main():
    out, svgs = sys.argv[1], sys.argv[2:]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(SLIDE_W), Emu(SLIDE_H)
    for s in svgs:
        n = convert(Path(s), prs, Path(s).stem)
        print(f"{Path(s).name}: {n} native shapes")
    prs.save(out)


if __name__ == "__main__":
    main()
